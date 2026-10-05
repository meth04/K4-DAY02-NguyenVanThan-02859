"""Package completed, real runs into the submission layout required by README.

No training or model inference: metrics are recomputed from existing CSV files.
Run: python code/finalize_submission.py --root local_runs/hour
"""
from __future__ import annotations

import argparse
import hashlib
import importlib.metadata
import json
import shutil
import sys
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO))
import eval as ev


def read(path):
    return json.loads(path.read_text(encoding="utf-8"))


def write(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, indent=2, ensure_ascii=False), encoding="utf-8")


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def copy(source, destination):
    destination.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, destination)


def table(frame):
    columns = list(frame.columns)
    lines = ["| " + " | ".join(columns) + " |", "|" + "|".join(["---"] * len(columns)) + "|"]
    for row in frame.itertuples(index=False, name=None):
        cells = [f"{v:.4f}" if isinstance(v, float) else str(v) for v in row]
        lines.append("| " + " | ".join(cells) + " |")
    return "\n".join(lines)


def group(root, exp, split):
    return ev.load_group(str(root / f"predictions/{exp}_seed*_{split}.csv"),
                         str(root / f"data/labels/{split}_subset0.csv"), ref_what=split)


def workbook(root, out, groups, timings):
    sheets = pd.read_excel(root / "results.xlsx", sheet_name=None)
    times = {r["exp_id"]: r for r in timings}
    inference = sheets["Inference"]
    for i, row in inference.iterrows():
        measurement = times[row.exp_id]
        for metric in ("p50", "p95", "p99"):
            inference.loc[i, f"{metric} (ms)"] = measurement[metric]
        inference.loc[i, "ảnh/s"] = 1000 / measurement["p50"]
        inference.loc[i, "chi phí so với I00"] = measurement["p50"] / times["I00"]["p50"]
        inference.loc[i, "ghi chú"] = (str(row["ghi chú"]) if pd.notna(row["ghi chú"]) else "") + "; timing: GPU forward + aggregate + softmax; 20 warmup/100 iterations; no decode/resize/H2D"
        if row.exp_id == "I05":
            inference.loc[i, "ghi chú"] = "B04+B03, both screening 1 epoch; unequal training budget vs F01; measured GPU aggregate timing"
    # The matched-seed T00 difference is zero: it is the reference itself.
    training = sheets["Training"]
    training.loc[training.exp_id == "T00", "Δ so với mốc"] = 0.0
    training.loc[training.exp_id == "T14", "trục thay đổi"] = "kết hợp color + label smoothing"
    for i, row in training.iterrows():
        pred = ev.read_pred(root / f"predictions/{row.exp_id}_seed{int(row.seed)}_val.csv")
        metric = ev.compute_metrics(pred.y_true, pred.y_pred, pred.probs)
        training.loc[i, "F1 Chinee apple val"] = metric["f1"][0]
        training.loc[i, "F1 Snake weed val"] = metric["f1"][7]
        reference = row["mốc so sánh"]
        training.loc[i, "nguồn log"] = f"evidence/runs/{row.exp_id}/seed{int(row.seed)}/history.csv"
        training.loc[i, "so sánh công bằng"] = "12 epochs, matched seed" if reference == "T00" else "1 epoch, seed0, vs TS00"
    finals, perclass = [], []
    for exp, (test, val) in groups.items():
        cfg = read(root / f"runs/{exp}/seed0/config.json")
        method = "hflip logits + temperature scaling" if exp == "F01" else "1-view, uncalibrated"
        description = f"DeiT tiny / {cfg['img_size']}px / full finetune / basic / CE / {cfg['epochs']} epochs / FP32 / {method}"
        for index, seed in enumerate(test.seeds):
            tm, vm = test.metrics[index], val.metrics[index]
            finals.append({"exp_id": exp, "cấu hình": description, "seed": seed,
                "macro-F1 val": vm["macro_f1"], "macro-F1 test": tm["macro_f1"],
                "top-1 test": tm["top1"], "ECE test": tm["ece"],
                "balanced accuracy test": tm["balanced_acc"], "NLL test": tm["nll"],
                "nguồn": f"predictions/{exp}_seed{seed}_test.csv"})
        finals.append({"exp_id": exp, "cấu hình": description, "seed": "mean ± std (0,1,2; ddof=1)",
            "macro-F1 val": ev.fmt(*val.summary["macro_f1"]),
            "macro-F1 test": ev.fmt(*test.summary["macro_f1"]),
            "top-1 test": ev.fmt(*test.summary["top1"]), "ECE test": ev.fmt(*test.summary["ece"]),
            "balanced accuracy test": ev.fmt(*test.summary["balanced_acc"]),
            "NLL test": ev.fmt(*test.summary["nll"]), "nguồn": f"evidence/evaluation/{exp}_summary.json"})
        names = ev.load_names(str(root / "data/labels/labels.csv"))
        for label, name in enumerate(names):
            pc = {"cấu hình": exp, "lớp": name, "số ảnh test": int(test.metrics[0]["support"][label])}
            for key, title in [("precision", "precision"), ("recall", "recall"), ("f1", "F1")]:
                pc[title] = float(test.summary[key][0][label])
                pc[title + " std"] = float(test.summary[key][1][label])
            perclass.append(pc)
    sheets["Final"] = pd.DataFrame(finals)
    sheets["PerClass"] = pd.DataFrame(perclass)
    sheets["Latency"] = pd.DataFrame([{
        "cấu hình": r["exp_id"], "GPU": r["gpu"], "dtype": r["dtype"], "batch": 1,
        "gộp BN": "không", "p50 (ms)": r["p50"], "p95 (ms)": r["p95"], "p99 (ms)": r["p99"],
        "ảnh/s": 1000/r["p50"], "warmup": r["warmup"], "n_iters": r["n"],
        "độ phân giải": r["img_size"], "phạm vi đo": r["scope"],
    } for r in timings])
    summary_rows = []
    for path in sorted((root / "runs").glob("*/seed0/summary.json")):
        result = read(path)
        exp = result["exp_id"]
        cfg = read(path.parent / "config.json")
        score = groups[exp][1].summary["macro_f1"] if exp in groups else (result["val_macro_f1"], np.nan)
        summary_rows.append({"exp_id": exp, "backbone": cfg["backbone"], "macro-F1 val": ev.fmt(*score),
            "val mean": score[0], "seed": "0,1,2" if exp in groups else "0", "epoch": cfg["epochs"],
            "công thức": f"{cfg['init']}/{cfg['aug']}/{cfg['loss']}",
            "p95 batch1 (ms)": times["F01" if exp == "F01" else "I00"]["p95"] if exp in groups else "",
            "ghi chú": "Selected on val; test unchanged" if exp == "F01" else "screening / different budget" if exp != "T00" else "matched baseline"})
    sheets["Summary"] = pd.DataFrame(summary_rows).sort_values("val mean", ascending=False).head(10).drop(columns="val mean")
    sheets["Robustness"] = pd.read_csv(root / "extras/robustness.csv")
    with pd.ExcelWriter(out / "results.xlsx", engine="openpyxl") as writer:
        for name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=name, index=False)
    book = load_workbook(out / "results.xlsx")
    for sheet in book:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="183B56")
            cell.font = Font(bold=True, color="FFFFFF")
            cell.alignment = Alignment(wrap_text=True)
        sheet.row_dimensions[1].height = 34
        for column in sheet.columns:
            width = min(50, max(14, max(len(str(cell.value or "")) for cell in column) + 2))
            sheet.column_dimensions[column[0].column_letter].width = width
        for row in list(sheet.rows)[1:]:
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "0.0000"
                if row[0].value == "F01":
                    cell.fill = PatternFill("solid", fgColor="DDF1E4")
        sheet.print_options.horizontalCentered = True
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
        sheet.print_title_rows = "1:1"
    book.save(out / "results.xlsx")
    return sheets


def figures(out, sheets, groups):
    for sheet, filename, x, y, title in [
        ("Backbones", "backbone_tradeoff.png", "độ trễ batch-1 (ms)", "macro-F1 val", "Backbones: screening, 1 epoch, seed0"),
        ("Inference", "inference_tradeoff.png", "p95 (ms)", "macro-F1 val", "Inference: validation F1 vs p95, GPU batch1 FP32")]:
        frame = sheets[sheet]
        fig, ax = plt.subplots(figsize=(8, 4.5))
        ax.scatter(frame[x], frame[y], s=70)
        for _, row in frame.iterrows():
            offset = (6, -14) if row.exp_id in ("I03", "I07") else (6, 8)
            ax.annotate(row.exp_id, (row[x], row[y]), xytext=offset, textcoords="offset points")
        ax.set(xlabel=x, ylabel="Validation macro-F1", title=title)
        ax.grid(alpha=.25)
        fig.tight_layout()
        fig.savefig(out / "curves" / filename, dpi=150)
        plt.close(fig)
    fig, ax = plt.subplots(figsize=(8, 4))
    frame = sheets["Robustness"]
    ax.bar(frame.condition, frame.macro_f1, color=["#27816D", "#7094B5", "#CB9158", "#B96666"])
    for index, row in frame.iterrows():
        ax.text(index, row.macro_f1 + .01, f"{row.macro_f1:.4f}", ha="center")
    ax.set(ylim=(0, 1), ylabel="Validation macro-F1", title="Synthetic domain shift: frozen F01 seed0 / 3501 validation images")
    fig.tight_layout()
    fig.savefig(out / "curves/robustness_val.png", dpi=150)
    plt.close(fig)
    # Each final seed has its own graph; alias T00 reuses exactly the same history.
    from train import plot_curves
    for exp in ("F01", "T00"):
        for seed in (0, 1, 2):
            history = pd.read_csv(out / f"evidence/runs/{exp}/seed{seed}/history.csv").to_dict("records")
            plot_curves(history, out / f"curves/{exp}_seed{seed}.png", f"{exp} seed{seed} / DeiT tiny / 128px")


def report(root, out, sheets, groups, timings):
    names = ev.load_names(str(root / "data/labels/labels.csv"))
    final, val = groups["F01"]
    base, _ = groups["T00"]
    uncal = group(root, "F01uncal", "test")
    delta = final.summary["macro_f1"][0] - base.summary["macro_f1"][0]
    noise = max(final.summary["macro_f1"][1], base.summary["macro_f1"][1])
    confusion = sum(m["confusion"] for m in final.metrics)
    pd.DataFrame(confusion, index=names, columns=names).to_csv(out / "evidence/evaluation/F01_confusion_sum.csv")
    off_diagonal = confusion.copy()
    np.fill_diagonal(off_diagonal, 0)
    pairs = sorted([(int(off_diagonal[i, j]), names[i], names[j]) for i in range(9) for j in range(9) if i != j], reverse=True)[:6]
    split = read(root / "runs/split_check.json")
    counts = pd.DataFrame({s: split["per_class"][s] for s in ["train", "val", "test"]})
    counts.index = names
    counts.insert(0, "Lớp", names)
    catalog = pd.read_csv(root / "data/labels/labels.csv").Label.value_counts().sort_index()
    counts["Catalog"] = catalog.values
    counts["Table 1 (tài liệu lab)"] = [1125,1064,1031,1022,1062,1009,1074,1016,9106]
    cfg = read(root / "runs/F01/seed0/config.json")
    environment = read(root / "environment.json")
    temperatures = [read(root / f"runs/F01/seed{s}/test_manifest.json")["temperature"] for s in range(3)]
    timing = next(t for t in timings if t["exp_id"] == "F01")
    train_frame = sheets["Training"].query("seed == 0 and exp_id != 'T00'")
    lines = ["# Lab Day 2 — Nguyễn Văn Thân (02859)", "", "## 1. Tóm tắt", "",
        "DeepWeeds gồm 9 lớp, đánh giá trên fold 0 gốc. Đã khảo sát 5 backbone, 3 trục training và một kết hợp; so sánh 4 phương pháp suy luận ngoài 1-view.",
        "Cấu hình chọn trên validation: **DeiT tiny ImageNet `fb_in1k`, 128 px, finetune toàn bộ, basic crop/flip, CE, 12 epoch, FP32; hflip gộp logits và temperature scaling**.",
        f"Test 3 seed (0,1,2): **top-1 {final.summary['top1'][0]*100:.2f}% ± {final.summary['top1'][1]*100:.2f} điểm phần trăm; macro-F1 {ev.fmt(*final.summary['macro_f1'])}**.",
        f"So với T00 + I00, Δ macro-F1 = **{delta:+.4f}**, nhỏ hơn std lớn hơn **{noise:.4f}**; chưa chứng minh cải thiện chắc chắn.",
        f"Độ trễ bổ sung của F01 gồm forward + gộp logits + softmax/temperature: p95 **{timing['p95']:.2f} ms**, batch 1 trên GTX 1650; chưa gồm decode/resize/chuyển ảnh lên GPU.",
        "", "## 2. Dữ liệu và thiết lập", "",
        "### 2.1. Fold và kiểm tra dữ liệu", "",
        "Giữ nguyên train/val/test CSV fold 0: **10.501 / 3.501 / 3.507** ảnh; giao từng cặp bằng 0, hợp đúng **17.509** ảnh. Chỉ train dùng để cập nhật trọng số. Mọi lựa chọn/checkpoint/T fit trên val; test chỉ chạy sau chốt lựa chọn, một lần cho mỗi cấu hình/seed (các view nằm trong cùng lần đánh giá).",
        table(counts.reset_index(drop=True)), "",
        "Catalog khớp Table 1. CSV fold gốc có đúng một khác biệt ở train: `20170714-110407-3.jpg` mang Label 0, catalog Label 1; giữ nguyên fold theo S1, không sửa/lọc ảnh. Vì thế tổng fold của lớp 0/1 là 1.126/1.063, còn catalog là 1.125/1.064. Kiểm tra đầy đủ: [split_check.json](evidence/runs/split_check.json).",
        "Negative chiếm 9.106/17.509 ≈ 52%; các lớp còn lại khoảng 1.009–1.125 ảnh. Do đó macro-F1 là tiêu chí chọn chính; top-1 phải đọc kèm F1/recall từng lớp.",
        "![Phân bố lớp](curves/eda_classes.png)", "![Ảnh mẫu](curves/eda_samples.png)",
        "", "### 2.2. Pipeline, phần cứng và tái lập", "",
        "Loss kiểm tra đầu 2,1104 so với ln(9)=2,1972; overfit batch nhỏ đạt **0,0131 sau 16 bước**. Sáu kiểm tra focal γ=0, LS ε=0, CutMix diện tích, Mixup nhãn, fuse Conv-BN và temperature đều pass. Đây là kiểm tra chức năng; không phải các thí nghiệm training độc lập.",
        "![Augmentation kiểm tra](curves/augmentation.png)",
        f"Phần cứng: Windows 11 / Ryzen 5 6600HS / RAM 16 GB / **{environment['gpu']} 4 GB**. Python 3.13.13, torch {environment['torch']}, torchvision {environment['torchvision']}, timm {environment['timm']}, NumPy {environment['numpy']}. Phiên bản phụ thuộc đầy đủ ở [environment.json](evidence/environment.json) và [requirements-local.txt](requirements-local.txt).",
        "Probe forward/backward/AdamW trên 5 mạng chọn **batch 64, FP32, NCHW, 2 workers, cuDNN benchmark tắt**, dùng chung cho mọi cấu hình. FP16 không nhanh hơn trên máy này; ConvNeXt FP16 batch64 có loss không hữu hạn nên loại. CUDA/cuDNN ở Colab khác máy này, vì vậy không quy mọi lỗi FIND về một nguyên nhân duy nhất.",
        f"Công thức nền: RandomResizedCrop(128, scale 0,6–1,0) + hflip; eval resize 146 rồi center crop 128, Normalize ImageNet. AdamW LR backbone {cfg['lr_backbone']}, head {cfg['lr_head']}, weight decay {cfg['weight_decay']}, warmup {cfg['warmup_epochs']} epoch + cosine. Backbone norm/bias không decay; nhóm head giữ decay theo bộ khung 3 nhóm. Không EMA/Mixup/sampler cân bằng. `eval()` và inference_mode dùng khi đánh giá.",
        "Khảo sát seed0, 1 epoch; chung kết 12 epoch cho đủ 3 seed. Lượt nghiên cứu chính mất **36,18 phút** (không tính cài đặt/tải weights/probe); các phép đo/giải thích bổ sung có log riêng. Không dùng benchmark này để hứa cùng thời gian trên Colab.",
        "", "## 3. So sánh backbone", "",
        table(sheets["Backbones"][["exp_id","backbone","tag trọng số","#params (M)","GMAC","macro-F1 val","top-1 val","train/epoch (s)","độ trễ batch-1 (ms)"]]),
        "![Đánh đổi backbone](curves/backbone_tradeoff.png)",
        "Cùng split/seed0, 128 px, 1 epoch, AdamW/crop/flip/CE/FP32; khác kiến trúc và tag pretrained công bố. DeiT tiny đạt F1 **0,7289** so với ConvNeXt atto **0,7128** (Δ=0,0161), thời gian epoch **54,4 so với 82,9 s**. ConvNeXt ít tham số hơn (3,38M so với 5,50M), nhưng DeiT hợp ngân sách local và có F1 val cao hơn nên đi tiếp. MobileNetV3 nhẹ nhất theo GMAC, nhanh nhất theo train/epoch; latency batch1 thực đo không nhất thiết thấp nhất.",
        "ResNet18/34 đạt F1 thấp sau chỉ 1 epoch; chưa đủ để kết luận họ ResNet kém khi hội tụ. GMAC là Conv2d + Linear, **chưa tính attention matmul**, nên số DeiT là cận dưới. Mỗi backbone mới có 1 seed và tag pretrained khác nhau: kết quả là lựa chọn thực dụng trong ngân sách, không phải kiểm định thuần kiến trúc.",
        "", "## 4. Ablation công thức huấn luyện", "",
        table(train_frame[["exp_id","trục thay đổi","macro-F1 val","Δ so với mốc","F1 Chinee apple val","F1 Snake weed val"]]),
        "TS00 là đúng bản chạy B04 dùng lại, cùng 1 epoch/seed0. T01 chỉ đổi finetune→frozen; T03 chỉ đổi basic→color jitter; T07 chỉ đổi CE→LS ε=0,1. Mỗi trục có hai giá trị và so với TS00, không so ablation 1 epoch với T00 12 epoch.",
        "Frozen giảm **0,3947 F1**: thích ứng backbone quan trọng trong chế độ khảo sát này. Color jitter giảm **0,0423**; giả thuyết thay đổi màu mạnh làm khó phân biệt lá khi chưa hội tụ, cần kiểm chứng thêm. LS giảm **0,0031**, không có std khảo sát nên không khẳng định nó có hại chắc chắn. T14 kết hợp color+LS đạt **0,6880**, tăng **0,0014** so với color đơn, nhưng vẫn thấp hơn TS00 **0,0409**; chưa có cộng dồn hữu ích. Hai thành phần này được thử để kiểm tra tương tác, không gọi là những yếu tố tốt nhất.",
        "Công thức thắng giữ nguyên basic + CE + full finetune. T00 và F01 dùng cùng trọng số train ở mỗi seed vì công thức training giống hệt; khác ở suy luận, có provenance trong summary. T00 không phải sáu lần train riêng. Thêm thời gian train từ 1 lên 12 epoch giúp F1 DeiT seed0 từ **0,7289 lên 0,9081** (khác lịch cosine/ngân sách; không quy toàn bộ cho một yếu tố).",
        "", "## 5. Suy luận và hiệu chuẩn", "",
        table(sheets["Inference"][["exp_id","phương pháp","K","macro-F1 val","top-1 val","ECE val","p50 (ms)","p95 (ms)","p99 (ms)"]]),
        "![Đánh đổi suy luận](curves/inference_tradeoff.png)",
        "Bốn phương pháp ngoài I00: hflip gộp xác suất (I01), hflip gộp logits (I03), ensemble B04+B03 (I05), temperature scaling (I07). Hflip logits tăng **0,0101 F1 val**; prob/logit chỉ chênh khoảng **0,000004**, không có bằng chứng logits tốt hơn chắc chắn. Quy tắc đã chốt trước test: giữ 1-view nếu gain≤0,001; lấy điểm val cao nhất nếu vượt ngưỡng. Temperature không đổi argmax/F1, giảm ECE I00 **0,0243→0,0054** với T≈1,3548.",
        "Ensemble 0,7542 F1 dùng hai checkpoint **khảo sát 1 epoch**, không tương đương ngân sách model F01 12 epoch. Không dùng số này để kết luận ensemble nói chung kém; nó cũng tốn hai mạng và thêm bộ nhớ.",
        "Bảng suy luận/latency cập nhật bằng phép đo bổ sung: batch1, GPU input 128 px đã sẵn, FP32, **20 warmup + 100 lượt**, synchronize trước/sau từng lượt, gồm forward và gộp/softmax. Latency khảo sát backbone và phép đo forward-only gốc (10 warmup/50 lượt) giữ ở evidence để truy ngược; không trộn hai phạm vi đo. Nghịch đảo p50 là thông lượng tuần tự batch1, không phải throughput batched.",
        "", "## 6. Chung kết và phân tích lỗi", "",
        ev.report_group("F01: hflip logits + temperature", final, names), "",
        ev.report_group("T00 + I00: baseline single-view", base, names), "",
        f"T fit riêng trên val theo seed0/1/2: **{', '.join(f'{t:.4f}' for t in temperatures)}**. Test ECE trước/sau: **{ev.fmt(*uncal.summary['ece'])} → {ev.fmt(*final.summary['ece'])}**. Gap F1 val/test = **{abs(val.summary['macro_f1'][0]-final.summary['macro_f1'][0]):.4f}**. Tất cả 6 manifest cuối có status complete và 3.507 ảnh; CSV uncal dùng cùng lượt logits đã lưu, không chạy test mới.",
        "![Ma trận nhầm lẫn test](curves/F01_confusion.png)",
        "Ma trận cộng 3 seed (cùng 3.507 ảnh được dự đoán ba lần, **không phải 10.521 ảnh độc lập**); hàng thật/cột dự đoán. Các nhầm lẫn lớn:",
        table(pd.DataFrame(pairs, columns=["Số lần qua 3 seed", "Thật", "Dự đoán"])),
        f"Cặp Chinee apple→Snake weed: **{int(confusion[0,7])}/{int(confusion[0].sum())}** lượt ({100*confusion[0,7]/confusion[0].sum():.2f}%); chiều ngược **{int(confusion[7,0])}/{int(confusion[7].sum())}** ({100*confusion[7,0]/confusion[7].sum():.2f}%). Đây không phải nguồn sai duy nhất; cần đọc cả nhầm với Negative và các loài lá tương tự.",
        "![Các ảnh test đoán sai seed0](curves/F01_errors.png)",
        "Trong các ảnh sai hiển thị, Chinee apple có thể nhỏ, nhiều lá/cành che lấp và bóng đổ; Snake weed có nền đất lớn hoặc lẫn nhiều thực vật. Hai ví dụ Chinee apple→Snake weed có ánh sáng/tông hồng khác ảnh lá xanh. Giả thuyết: giảm xuống 128 px làm mất chi tiết hình thái và tín hiệu nền/màu có thể chi phối. Ảnh minh họa không đủ xác nhận nhân quả; không sửa nhãn hoặc cấu hình dựa vào quan sát test.",
        "### 6.1. Đường cong training", "",
        "[F01 seed0](curves/F01_seed0.png), [seed1](curves/F01_seed1.png), [seed2](curves/F01_seed2.png); tất cả B/T/TS có ảnh và CSV riêng. Seed0: train loss giảm **0,8791→0,0438**, val loss **0,5563→0,2257**; F1 val dao động (epoch5 0,8801 xuống epoch6 0,8632) rồi tốt nhất epoch11 0,9081, epoch12 0,9074. Khoảng train/val loss cuối khá lớn gợi ý overfit; chọn checkpoint bằng macro-F1 val, không lấy epoch cuối mặc định. Seed1/2 chọn epoch12/10. Đồ thị 1 điểm của khảo sát chỉ chứng minh lần chạy, không cho phép kết luận tốc độ hội tụ.",
        "", "## 7. Kết luận và khuyến nghị", "",
        f"F01 đạt macro-F1 test {ev.fmt(*final.summary['macro_f1'])}; Δ so mốc {delta:+.4f} < std {noise:.4f}, nên mức tăng quan sát còn trong nhiễu. Trong khảo sát có kiểm soát, finetune backbone có tác động lớn nhất (**+0,3947 so frozen**); khác biệt backbone tốt nhất với thứ hai nhỏ hơn (**0,0161**), TTA cải thiện val **0,0101**. Tăng ngân sách train có thay đổi lớn (**+0,1792**, kèm đổi lịch LR), nhưng các phép so không cùng điều kiện nên không xếp hạng đóng góp nhân quả chung.",
        f"Với GPU khai báo và ngân sách **30–100 ms cho model**, chọn F01 hflip+T: p95 {timing['p95']:.2f} ms và F1 test {final.summary['macro_f1'][0]:.4f}. Nếu ngân sách end-to-end chặt hoặc GPU yếu hơn, T00 1-view có F1 test {base.summary['macro_f1'][0]:.4f}, ít chi phí hơn; phải đo cả camera/decode/resize/H2D trước triển khai. Không suy độ trễ GTX 1650 sang Jetson/robot khác. F01 phù hợp ngoại tuyến hoặc hệ thống đủ ngân sách; độ tin cậy nên kèm ngưỡng từ validation riêng của môi trường triển khai.",
        "", "## 8. Thí nghiệm bổ sung trên validation", "",
        "### 8.1. Lệch phân phối nhân tạo", "",
        "Cố định F01 seed0 và T đã fit trên val sạch; **không train/thích ứng/refit T**, không dùng test. Đánh giá toàn bộ 3.501 ảnh val ở 4 điều kiện: sạch; brightness×0,55; Gaussian blur radius1,5 ở ảnh gốc 256px; Gaussian noise σ20/255, seed20261005+row. Hflip logits giữ nguyên. Đây là kiểm tra độ bền mô tả, không chọn lại model.",
        table(sheets["Robustness"][["condition","n","macro_f1","top1","ece_before","ece_after"]]),
        "![Lệch miền validation](curves/robustness_val.png)",
        "F1 giảm từ **0,9182** xuống **0,8167** khi tối (−0,1015), **0,6101** khi mờ (−0,3081), **0,8763** khi nhiễu (−0,0419); điều kiện mờ ảnh hưởng lớn nhất trong ba mức đã chọn. T sạch giảm ECE cả ba miền, nhưng ECE sau hiệu chuẩn khi mờ vẫn **0,1410**, lớn hơn **0,0076** ở sạch. Ánh sáng/chi tiết/nhiễu có thể làm giảm chất lượng dù cùng loài/ảnh. Chỉ 1 seed và corruptions nhân tạo, không thay thế tập đánh giá theo địa điểm/mùa hay đảm bảo calibration ở mọi miền lệch. Xem CSV đầy đủ và protocol ở `evidence/extras/`.",
        "### 8.2. Giải thích bằng Grad-CAM", "",
        "![Grad-CAM ảnh validation sai](curves/gradcam_val.png)",
        "Sáu lỗi val (3 Chinee apple, 3 Snake weed) chọn theo thứ tự file, không cherry-pick theo heatmap. Grad-CAM cho lớp dự đoán **single-view**, hook `blocks[-1].norm1` trước attention cuối, bỏ CLS, reshape 8×8 patch, trọng số bằng gradient trung bình, ReLU và normalize. Đã kiểm tra gradient hữu hạn/khác0; label single-view và TTA ghi riêng. Bản đồ mức patch thô, chỉ gợi ý vùng nhạy với score; không phải mặt nạ cỏ/định vị đúng và không chứng minh model đã hiểu hình thái.",
        "Ví dụ `20170410-153511-0.jpg` (Chinee apple→Parthenium) có phản hồi lớn ở nhóm lá nửa dưới và mép phải; `20170718-100650-2.jpg` (→Rubber vine) phản hồi rải trên cả lá/cành; `20170711-114614-0.jpg` (Snake weed→Chinee apple) nổi vùng lá phía trên trái. Vùng chú ý rộng trong nền nhiều thực vật phù hợp giả thuyết thiếu chi tiết phân biệt ở128px, nhưng không chứng minh nguyên nhân. Không dùng quan sát này để đổi cấu hình test.",
        "", "## 9. Hạn chế và phụ lục", "",
        "Một fold, 3 seed chung kết nhưng chỉ 1 seed khảo sát; 1 epoch sàng lọc có thể đổi thứ hạng khi hội tụ. Các ablation color/LS chưa có std. Không khảo sát scratch/CutMix/focal trong lượt chính, chỉ có unit checks chức năng. Không chạy nhiều fold, distillation, adaptation hoặc ONNX; không nhận các điểm thưởng này.",
        "Fold ngẫu nhiên không theo địa điểm có thể cho test lạc quan; chưa kiểm chứng dữ liệu thật theo mùa/camera. Pretrained tags khác nhau và giảm128px khiến so với bài báo chỉ tham khảo. Recall Chinee apple 79,2% còn dưới mốc 88,5%; Snake weed 88,9% gần mốc 88,8% của tài liệu, điều kiện bài báo khoảng100epoch và 5fold khác thí nghiệm này.",
        "Các sự cố đã sửa trước lượt chính: khác biệt catalog/fold nhãn, numpy bool JSON và CUDA/cuDNN/precision. Probe FP16 không hữu hạn được ghi, không thay bằng số đẹp. Sau khi mở test không có huấn luyện/tuning/test mới; chỉ đọc CSV cũ để hoàn thiện sản phẩm và thêm thí nghiệm val riêng.",
        "`eval.py` gốc tự chấm **10/20 cho mục I**, chưa phải điểm toàn bài. A–H và điểm thưởng cần giảng viên đánh giá minh chứng; không cam kết điểm tối đa. Muốn nâng model ở nghiên cứu tiếp theo cần split/test độc lập mới, không tiếp tục chọn cấu hình trên test đã mở.",
        "Danh mục cấu hình/log: [evidence/runs/](evidence/runs/); workbook [results.xlsx](results.xlsx); hướng dẫn tái lập [README.md](README.md); đối chiếu rubric [RUBRIC_EVIDENCE.md](RUBRIC_EVIDENCE.md). Mọi file có SHA-256 trong [manifest.json](manifest.json). Tài liệu gốc: [README lớp](evidence/lab_docs/README.md), [GUIDE](evidence/lab_docs/GUIDE.md), [RUBRIC](evidence/lab_docs/RUBRIC.md).",
        ""]
    (out / "report.md").write_text("\n\n".join(lines), encoding="utf-8")


def readme(root, out, timing):
    text = f"""# Bài nộp Lab Day 2 — Nguyễn Văn Thân (02859)

Kết quả thực nghiệm local: **accuracy 93,67% ± 0,30 điểm phần trăm; macro-F1 0,9139 ± 0,0044**, 3 seed trên toàn bộ test fold0. Báo cáo đã hoàn thiện từ lần train 36,18 phút; phần bổ sung chỉ validation và đo latency, không train/test lại.

- [Báo cáo](report.md), [workbook 8 sheet](results.xlsx), [biểu đồ](curves/), [dự đoán](predictions/), [minh chứng rubric](RUBRIC_EVIDENCE.md).
- [Mở notebook trên Google Colab](https://colab.research.google.com/github/meth04/K4-DAY02-NguyenVanThan-02859/blob/main/code/lab_day2.ipynb).
- Mã training thực dùng tại `9747314cb5aa60fc727fefaa089aacc3de544d15`; snapshot trong `code/`, hash core trong `evidence/runs/source_manifest.json`. Notebook hiện tải bản sửa renderer `cd6b0eb2aafe44b30cead7eb8fb20a0885c30c95`; không thay core training.

## Kiểm tra điểm từ CSV, không cần GPU hoặc checkpoint

Từ thư mục bài nộp này, cài `numpy` và `pandas` rồi chạy:

```powershell
python eval.py score --pred 'predictions/F01_seed*_test.csv' --test-csv evidence/labels/test_subset0.csv --labels evidence/labels/labels.csv --tag F01 --out verification
python eval.py grade --final 'predictions/F01_seed*_test.csv' --baseline 'predictions/T00_seed*_test.csv' --uncal 'predictions/F01uncal_seed*_test.csv' --final-val 'predictions/F01_seed*_val.csv' --test-csv evidence/labels/test_subset0.csv --val-csv evidence/labels/val_subset0.csv --labels evidence/labels/labels.csv --latency-p95-ms {timing['p95']:.10f} --out verification
```

`eval.py` nguyên bản, hash kiểm tra trong manifest. `grade` hiện cho **10/20 riêng mục I**, không tự chấm A–H hoặc điểm thưởng. `evidence/labels/` là bốn CSV metadata gốc để kiểm tra fold offline; không có ảnh dataset. Prediction test, val và uncal đều được giữ lại; bản uncal là cùng lượt dự đoán chưa temperature scaling.

## Chạy lại training

Từ root repository, xem [hướng dẫn local](https://github.com/meth04/K4-DAY02-NguyenVanThan-02859/blob/main/SUBMISSION_README.md). Môi trường đo: Windows11, Python3.13.13, torch2.7.1+cu126, torchvision0.22.1+cu126, timm1.0.30, GTX1650 4GB. `requirements-local.txt` ghi phiên bản runtime còn lại.

```powershell
python -m venv --system-site-packages .venv-gpu
.\\.venv-gpu\\Scripts\\python.exe -m pip install torch==2.7.1 torchvision==0.22.1 --index-url https://download.pytorch.org/whl/cu126
.\\.venv-gpu\\Scripts\\python.exe -m pip install -r requirements.txt
.\\.venv-gpu\\Scripts\\python.exe -X utf8 -u code/run_local.py --root local_runs/reproduction --minutes 60
```

Dùng thư mục mới cho nghiên cứu mới. CLI probe GPU và chọn batch/precision theo tốc độ, epoch chung kết theo ngân sách; có thể khác cấu hình bản nộp trên phần cứng khác. Để tái lập cấu hình đã nộp, dùng settings trong `evidence/environment.json` và config từng seed: 128px, batch64, FP32/NCHW, workers2, cudnn benchmark off; khảo sát1epoch, final12epoch, seed0/1/2. Test một lần sau khi chốt trên val; không sửa cấu hình dựa vào test.

Colab: chọn GPU, PROFILE=hour (128px), SAVE_TO_DRIVE=True, chạy cell theo thứ tự. Notebook hour mặc định final2epoch; thêm `lab.budget["final"] = 12` ngay sau khi tạo `lab`, trước phần backbone, nếu cần tái lập số epoch local. Để giữ precision/batch giống bản nộp, cập nhật `settings` trước khi tạo lab: `batch_size=64, amp=False, channels_last=False, num_workers=2, cudnn_benchmark=False`. Có thể cần giảm batch theo VRAM GPU được cấp, khi đó ghi rõ khác biệt. Giữ SESSION riêng cho mỗi nghiên cứu để không trộn nguồn/checkpoint. Thời gian và kết quả không thể bảo đảm giống trên phần cứng khác.

## Tái lập minh chứng bổ sung và đóng gói

Sau khi có checkpoint/dataset local cùng nghiên cứu:

```powershell
.\\.venv-gpu\\Scripts\\python.exe -X utf8 -u code/submission_extras.py --root local_runs/hour
.\\.venv-gpu\\Scripts\\python.exe -X utf8 code/finalize_submission.py --root local_runs/hour
```

Extras chỉ đo GPU latency và validation (không train/test). Packaging chỉ đọc CSV/log đã có; không gọi model. Dataset images.zip/ảnh/checkpoint **không commit** theo README lớp. Các checkpoint hiện nằm ở `local_runs/hour/runs/` trên máy chủ bài nộp; không có link checkpoint công khai. Người chấm kiểm tra toàn bộ metric test trực tiếp từ CSV, hoặc train lại bằng code/notebook.

Các phép đo bổ sung: 20 warmup/100 lượt/synchronize, batch1 FP32, gồm forward+aggregate+softmax; không gồm camera/decode/resize/H2D. Kết quả mới và forward-only gốc đều có evidence, tránh đánh đồng latency model với end-to-end robot.
"""
    (out / "README.md").write_text(text, encoding="utf-8")


def rubric_evidence(out):
    text = """# Đối chiếu rubric và minh chứng

Các mục dưới đây là vị trí minh chứng, **không phải cam kết điểm do sinh viên tự chấm**. Tham chiếu [RUBRIC lớp](evidence/lab_docs/RUBRIC.md). Mọi lựa chọn đã chốt trước test; báo cáo giữ các kết quả yếu/thí nghiệm không có lợi.

| Mục | Minh chứng | Giới hạn cần người chấm xem |
|---|---|---|
| P1 — 6 sản phẩm | README.md, results.xlsx, report.md, curves/, code/, predictions/ | Có đủ trong thư mục bài nộp |
| P2 — số liệu thật | [training.log](evidence/training.log), [runs](evidence/runs/), [manifest](manifest.json) | Mỗi exp/seed có config/history/summary; T00 và TS00 là alias có provenance |
| P3 — fold 0 | [split_check.json](evidence/runs/split_check.json), [CSV gốc](evidence/labels/) | Giữ nguyên khác biệt 1 nhãn catalog/fold ở train; không tự chia/sửa CSV |
| P4 — dự đoán đủ seed | [predictions](predictions/), [grade_I.json](evidence/evaluation/grade_I.json) | F01/T00 đều đủ seed0/1/2, 3507 ảnh test/seed; có val và uncal |
| A — pipeline (12) | Report §2, EDA/augmentation PNG, [pipeline_checks.json](evidence/runs/pipeline_checks.json), environment/config | Loss đầu/overfit/checks có số thật; test manifests complete |
| B — backbone (12) | Backbones sheet, report §3, B01–B05 curves/config/summary | 5 mạng, ResNet/ConvNeXt/transformer/MobileNet; khảo sát chỉ 1epoch/1seed; GMAC transformer là cận dưới |
| C — training (16) | Training sheet, report §4, TS00/T01/T03/T07/T14 | 3 trục × 2 giá trị; các thử đơn đổi 1 yếu tố; kết hợp color+LS không thắng; chưa có std ablation |
| D — inference (12) | Inference/Latency sheets, report §5, [timings](evidence/extras/inference_timings.json), tradeoff PNG | 4 phương pháp ngoài I00; 20warmup/100 synchronized iterations; ensemble chỉ 1epoch nên không so ngang ngân sách final |
| E — workbook (8) | [results.xlsx](results.xlsx) | 7 sheet bắt buộc + Robustness; Final có từng seed và dòng mean±std; header/freeze/filter/units |
| F — curves (4) | [curves](curves/), evidence/runs/*/seed*/history.csv | Mỗi B/T/F có ảnh; F01/T00 còn có từng seed; 1epoch screening chỉ một điểm |
| G — báo cáo (12) | [report.md](report.md) §1–9, confusion/errors/Grad-CAM | Có kết luận, Δ so std, F1/recall lớp khó, lỗi ảnh, hạn chế; không còn TODO |
| H — tái lập (4) | [README](README.md), [code](code/), [requirements-local](requirements-local.txt), eval.py gốc | Code hoàn thiện trong code/, không sửa starter hay evaluator lớp; notebook Colab và lệnh kiểm tra offline |
| I — model (20) | [grade_I.json](evidence/evaluation/grade_I.json), Final/PerClass, predictions | eval.py tự chấm 10/20 (ngưỡng tạm thời); không phải tổng bài |

## Phần bổ sung xin xem xét điểm thưởng

| Thí nghiệm thật | Minh chứng | Mức trong rubric / giới hạn |
|---|---|---|
| Grad-CAM giải thích lỗi | Report §8.2, [ảnh](curves/gradcam_val.png), [thông tin](evidence/extras/gradcam_val.json), code/submission_extras.py | Mục +1; 6 lỗi validation, single-view predicted class, không dùng test để chọn mẫu |
| Phân tích lệch phân phối | Report §8.1, Robustness sheet, [protocol](evidence/extras/robustness_protocol.json), 4 CSV val đầy đủ | Mục +2; 3501 ảnh/điều kiện, seed0, tối/mờ/nhiễu, ECE trước/sau T cố định từ sạch |

Không nhận điểm nhiều fold/DINOv2/distillation/adaptation/ONNX vì chưa có thực nghiệm. Giảng viên xác nhận điểm A–H và thưởng; số epoch khảo sát ngắn và những thiếu sót nêu trong báo cáo có thể ảnh hưởng điểm.
"""
    (out / "RUBRIC_EVIDENCE.md").write_text(text, encoding="utf-8")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=REPO / "local_runs/hour")
    parser.add_argument("--out", type=Path, default=REPO / "submissions/02859_nguyen_van_than")
    args = parser.parse_args()
    root, out = args.root.resolve(), args.out.resolve()
    if not out.is_relative_to(REPO / "submissions"):
        raise ValueError("Submission output must stay inside this repository's submissions directory.")
    out.mkdir(parents=True, exist_ok=True)
    assert read(root / "local_complete.json")["status"] == "complete"
    assert sha(root / "eval.py") == sha(REPO / "eval.py"), "eval.py differs from the original"
    for name, expected in read(root / "runs/source_manifest.json").items():
        assert sha(root / "code" / name) == expected, f"Training snapshot changed: {name}"
    # Confirm every existing prediction is byte-identical to the completed study
    # archive created before documentation/bonus work began.
    from zipfile import ZipFile
    with ZipFile(root / "deepweeds_submission.zip") as original:
        for path in (root / "predictions").glob("*.csv"):
            assert original.read("predictions/" + path.name) == path.read_bytes(), path
    # Copy only explicit small deliverables, never checkpoint/data archives.
    for folder in ("predictions", "curves"):
        for path in sorted((root / folder).glob("*")):
            if path.is_file() and path.suffix in {".csv", ".png"}:
                copy(path, out / folder / path.name)
    for path in (root / "runs").rglob("*"):
        if path.is_file() and path.suffix in {".json", ".csv", ".txt"}:
            copy(path, out / "evidence/runs" / path.relative_to(root / "runs"))
    for path in (root / "data/labels").glob("*.csv"):
        copy(path, out / "evidence/labels" / path.name)
    for name in ("README.md", "GUIDE.md", "RUBRIC.md"):
        copy(REPO / name, out / "evidence/lab_docs" / name)
    for path in (root / "extras").glob("*"):
        copy(path, out / "evidence/extras" / path.name)
    copy(root / "extras/gradcam_val.png", out / "curves/gradcam_val.png")
    for name in ("environment.json", "selection.json", "local_plan.json", "cuda_probe.json", "local_complete.json", "training.log"):
        copy(root / name, out / "evidence" / name)
    for path in (root / "code").glob("*.py"):
        copy(path, out / "code" / path.name)
    for name in ("submission_extras.py", "finalize_submission.py", "verify_submission.py", "make_results.py", "make_report.py", "lab_day2.ipynb"):
        copy(REPO / "code" / name, out / "code" / name)
    copy(REPO / "eval.py", out / "eval.py")
    copy(REPO / "requirements.txt", out / "requirements.txt")
    dependencies = ["numpy", "pandas", "Pillow", "scikit-learn", "matplotlib", "openpyxl", "timm"]
    (out / "requirements-local.txt").write_text("# Install CUDA torch/torchvision separately as documented in README.\n" + "\n".join(f"{name}=={importlib.metadata.version(name)}" for name in dependencies) + "\n", encoding="utf-8")
    groups = {exp: (group(root, exp, "test"), group(root, exp, "val")) for exp in ("F01", "T00")}
    for exp, (test, val) in groups.items():
        assert test.seeds == val.seeds == [0,1,2]
        assert all(metric["n"] == 3507 for metric in test.metrics)
        ev.save_group(out / "evidence/evaluation", exp, test, ev.load_names(str(root / "data/labels/labels.csv")))
    timings = read(root / "extras/inference_timings.json")
    sheets = workbook(root, out, groups, timings)
    figures(out, sheets, groups)
    report(root, out, sheets, groups, timings)
    timing = next(row for row in timings if row["exp_id"] == "F01")
    readme(root, out, timing)
    rubric_evidence(out)
    result = ev.main(["grade", "--final", str(root / "predictions/F01_seed*_test.csv"),
        "--baseline", str(root / "predictions/T00_seed*_test.csv"), "--uncal", str(root / "predictions/F01uncal_seed*_test.csv"),
        "--final-val", str(root / "predictions/F01_seed*_val.csv"), "--test-csv", str(root / "data/labels/test_subset0.csv"),
        "--val-csv", str(root / "data/labels/val_subset0.csv"), "--labels", str(root / "data/labels/labels.csv"),
        "--latency-p95-ms", str(timing["p95"]), "--out", str(out / "evidence/evaluation")])
    assert result == 0
    # Curated report and required workbook must have no unfilled draft text.
    assert "TODO" not in (out / "report.md").read_text(encoding="utf-8")
    assert len(sheets["Backbones"]) >= 5 and len(sheets["Inference"]) >= 5
    assert len(sheets["Final"]) == 8 and len(sheets["PerClass"]) == 18
    from zipfile import ZipFile, ZIP_DEFLATED
    manifest = {}
    for path in sorted(out.rglob("*")):
        if not path.is_file() or path.name == "manifest.json":
            continue
        assert path.suffix not in {".pt", ".pth", ".zip", ".npy", ".pyc", ".jpg"}, path
        assert path.stat().st_size < 25_000_000, path
        manifest[path.relative_to(out).as_posix()] = dict(bytes=path.stat().st_size, sha256=sha(path))
    write(out / "manifest.json", manifest)
    archive = root / "deepweeds_submission_final.zip"
    with ZipFile(archive, "w", compression=ZIP_DEFLATED) as bundle:
        for path in sorted(out.rglob("*")):
            if path.is_file():
                bundle.write(path, path.relative_to(out).as_posix())
    with ZipFile(archive) as bundle:
        assert len(bundle.namelist()) == len(set(bundle.namelist())) and bundle.testzip() is None
    print(f"Submission: {out}\nFiles: {len(manifest)+1}\nFinal archive: {archive}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
